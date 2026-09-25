"""PII Redaction Tool - hybrid regex + document-specific entity dictionaries.

Usage: python pii_redaction_v2.py INPUT.docx OUTPUT.docx
"""
from __future__ import annotations
import re
import sys
from docx import Document


PERSONS = [
    "Kushal Subbayya Hegde", "Pushpa Kushal Hegde", "Rajesh Kushal Hegde",
    "Rohit Kushal Hegde", "Rakhi Girija Shetty", "Dinesh Hirachand Munot",
    "Ajay Shriram Patil", "Ram Kumar Tiwari", "Indu Jacob", "Sarthak Malvadkar",
    "Sandesh Bhagwat", "Amod Joshi", "Ganesh Prasad", "Lalit Muljibhai Sarvaiya",
    "Lokesh Shah", "Soumavo Sarkar", "Kishan Rastogi", "Abhijit Diwan",
    "Shanti Gopalkrishnan", "Prakash Boricha", "Sheetal Parab", "Eric Bacha",
    "Sachin Gawade", "Pravin Teli", "Siddharth Jadhav", "Tushar Gavankar",
    "Varun Badai", "Hitesh Ramani", "Chitra Raste", "Sharmila Joshi",
    "Cherag Gyara", "Manisha Shukla", "Tushar Wakhele", "Ashish Mathew Pulloor",
    "Anand Soni", "Parag Pansare", "Sangeeta Ramprasad Rai",
]

COMPANIES = [
    "Bhandary Metal Extrusion Private Limited", "KSH International Private Limited",
    "KSH International Limited", "Nuvama Wealth Management Limited",
    "ICICI Securities Limited", "MUFG Intime India Private Limited",
    "Link Intime India Private Limited", "HDFC Bank Limited", "HDFC Limited",
    "ICICI Bank Limited", "IndusInd Bank Limited", "State Bank of India",
    "The Federal Bank Limited", "Bajaj Finance Limited", "Citibank N.A.",
    "Export-Import Bank of India", "National Securities Depository Limited",
    "Bharat Bijlee Limited", "CARE Ratings Limited", "CARE Analytics and Advisory Private Limited",
    "Care Ratings Limited", "Elantas Beck India Limited", "Hindalco Industries Limited",
    "Nidec Industrial Automation India Private Limited", "Precision Wires India Limited",
    "Savli Copper Products Private Limited", "Vedanta Limited", "BSE Limited",
    "Waterloo Motors Private Limited", "KSH Project Management Services Private Limited",
    "KSH Infra Park 5 Private Limited", "KSH Infra Park VI Private Limited",
    "KSH Distriparks Private Limited", "KSH Integrated Logistics Private Limited",
    "Kushal Motors and Electricals Private Limited", "Waterloo Industrial Park I Private Limited",
    "Waterloo Industrial Park II Private Limited", "Waterloo Industrial Park III Private Limited",
    "Waterloo Industrial Park IV Private Limited", "Waterloo Industrial Park V Private Limited",
    "Waterloo Industrial Park VI Private Limited", "Waterloo Industrial Park VIII Private Limited",
    "Waterloo Industrial Park IX Private Limited", "Waterloo Industrial Park IX A Private Limited",
    "Waterloo Industrial Park IX B Private Limited", "KSH Infra Park IV Private Limited",
    "Kirtane & Pandit LLP", "Trilegal", "Malabar India Fund Limited",
]

# High-confidence exact addresses observed in the supplied prospectus.
KNOWN_ADDRESSES = [
    "11/3, 11/4 and 11/5, Village Birdewadi, Chakan Taluka - Khed, Pune – 410 501, Maharashtra, India",
    "201, Tower 2, Montreal Business Centre, Off Pallod Farms, Baner, Pune – 411 045, Maharashtra, India",
    "S. no. 245/ 104, Pushpakamal, Deccan Gymkhana Society, lane no. 3 Prabhat Road, opposite PYC basketball court, Deccan Gymkhana, Pune – 411 004 Maharashtra, India",
    "12 Buena Monte, NCL co-operative housing society, Panchvati, Pashan, Pune – 411 008, Maharashtra, India",
    "Pushpakamal Apartment, Flat – 1, S. no. 245/ 104, Prabhat Road Lane no. 3, Shivaji Nagar, Deccan Gymkhana, Pune – 411 004, Maharashtra, India",
    "3 Prabhat Road, opposite PYC basketball court, Erandawane, Deccan Gymkhana, Pune – 411 004 Maharashtra, India",
    "Pratik Bunglow, Senapati Bapat Road, behind Sahara Hotel, Shivajinagar, Model Colony, Pune – 411 016, Maharashtra, India",
    "602, Gopalkrupa Apartment, Bhonde colony, Prabhat Road, Erandawane, Pune – 411 004, Maharashtra, India",
    "A-259, JK Road, Minal Residency, Huzur, Govindpura, Bhopal – 462 023, Madhya Pradesh, India",
    "A29, Abhimanshree Society, Pashan Road, Pune – 411 008, Maharashtra, India",
    "Plot No. J-25, Taloja Industrial Area, Village Padghe, Taluka Panvel, Raigad – 410 208, Maharashtra, India",
    "Plot No. 5, Chakan Industrial Area, Phase II, Village Khalumbre, Taluka Khed, Pune – 410 501, Maharashtra, India",
    "Plot No. F-223, Supa Parner Industrial Park, Mauje Palve Khurd, Taluka Parner, Dist – Ahmednagar, Maharashtra – 414 301",
    "801 - 804, Wing A, Building No 3, Inspire BKC, G Block, Bandra Kurla Complex, Bandra East, Mumbai 400051, Maharashtra, India",
    "ICICI Venture House, Appasaheb Marathe Marg, Prabhadevi, Mumbai 400025, Maharashtra, India",
    "C-101, Embassy 247, 1st Floor, L B S Marg, Vikhroli (West), Mumbai 400083, (Maharashtra), India",
    "One World Centre, 10th Floor, Tower 2A & 2B, Senapati Bapat Marg, Lower Parel (West), Mumbai – 400 013, Maharashtra, India",
    "HDFC Bank Limited, FIG-OPS Department – Lodha I Think Techno Campus, O-3 Level, Next to Kanjurmarg Railway Station, Kanjurmarg (East) Mumbai – 400042, Maharashtra, India",
    "163, 5th Floor, H.T.Parekh Marg, Backbay Reclamation Churchgate, Mumbai – 400020",
    "PCNTDA Green Building Block A 1st and 2nd floor Near Akurdi Railway Station Akurdi, Pune – 411 044 Maharashtra, India",
    "5th Floor, Wing A, Gopal House, S. No. 127/1B/1, Plot A1, Opp Harshal Hall, Kothrud, Pune – 411 038, Maharashtra, India",
    "Hingne Tare & Associates Flat No. 102, Sai Complex Shaniwar Peth, Pune – 411 030 Maharashtra, India",
    "8th Floor, Onyx Tower, North Main Road, Koregaon Park, Pune – 411 001 Maharashtra, India",
    "No. 401, 401(A), 401(B) & 402, 402(A), 402(B), 4th Floor, Signature Building, Bhandarkar road Shivaji Nagar, Pune – 411 004 Maharashtra, India",
    "2401 Gen Thimmayya Road, Cantonment Pune – 411 001 Maharashtra, India",
    "ICICI Bank, CBG, 3rd Floor, 362, Satguru House, Next to Tanishq Showroom, CTS No. 30, Bund Garden Road, Pune – 411 001 Maharashtra, India",
    "5th Floor, Marathon IT Park, Bund Garden Road, Pune – 411 001 Maharashtra, India",
    "Tara Chambers, Mumbai-Pune Road, Wakdewadi Pune – 411 003 Maharashtra, India",
    "Ground Floor, Kubera Chambers, Opp. Sancheti Hospital Shivajinagar, Pune – 411 005 Maharashtra, India",
    "The Capital, Unit no. 1601, B-wing, BKC, Mumbai, Maharashtra, India",
    "Gat No. 11/3, 11/4, 11/5, Village Birdewadi",
    "Gat No. 11/3, 11/4, 11/5, Village Birdewadi Taluka Khed, District Pune – 410 501 Maharashtra, India",
]

EMAIL_RE = re.compile(r"(?i)\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b")
SSN_RE = re.compile(r"\b\d{3}-\d{2}-\d{4}\b")
CC_RE = re.compile(r"(?<!\d)(?:\d[ -]?){13,19}(?!\d)")
IP_RE = re.compile(r"\b(?:(?:25[0-5]|2[0-4]\d|1?\d?\d)\.){3}(?:25[0-5]|2[0-4]\d|1?\d?\d)\b")
DOB_RE = re.compile(r"(?i)\b(?:date\s+of\s+birth|dob|birth\s+date)\s*[:\-]?\s*(\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|\d{1,2}\s+(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\s+\d{4})")
PHONE_RE = re.compile(r"(?<!\d)(?:\+?\s*91[\s-]?)?(?:\(?0?\d{2,4}\)?[\s-]?)?(?:\d{3,5}[\s-]\d{4}|\d{8}|\d{10})(?!\d)")
GENERIC_ADDRESS_RE = re.compile(r"(?i)(?<![A-Za-z0-9])(?:\d{1,5}/\s*\d{1,5}|[A-Z]-\d+|C-\d+|S\.?\s*no\.?|Gat\s+No\.?|Plot\s+No\.?|A\d{1,3})[^;\n]{0,220}?\b\d{3}\s*\d{3}\b")


def valid_phone(s: str) -> bool:
    digits = re.sub(r"\D", "", s)
    if len(digits) not in (10, 11, 12):
        return False
    if len(digits) == 10 and re.fullmatch(r"(?:19|20)\d{2}\d{6}", digits):
        return False
    return bool(re.search(r"\+?\s*91", s) or re.search(r"\b0\d{2,4}[\s-]\d{7,8}\b", s))

class Redactor:
    def __init__(self):
        self.mapping = {}
        self.counters = {k:0 for k in ["name","email","phone","company","address","ssn","credit_card","dob","ip"]}
        self.stats = {k:0 for k in self.counters}

    def replacement(self, kind, original):
        norm = re.sub(r"\s+", " ", original).strip().casefold() if kind in {"name","email","company","address"} else original
        key=(kind,norm)
        if key not in self.mapping:
            self.counters[kind]+=1
            i=self.counters[kind]
            if kind=="name": val=["John Carter","Emily Wilson","Daniel Parker","Sophia Miller","Michael Brown","Olivia Taylor","James Anderson","Emma Davis","Robert Martin","Grace Thompson","William Clark","Ava Lewis"][i%12]
            elif kind=="email": val=f"person{i:03d}@example.com"
            elif kind=="phone": val=f"+91 90000 {10000+i:05d}"
            elif kind=="company": val=f"Example Company {i:02d} Private Limited"
            elif kind=="address": val=f"{100+i}, Example Road, Pune – 411 {100+i:03d}, Maharashtra, India"
            elif kind=="ssn": val=f"999-99-{1000+i:04d}"
            elif kind=="credit_card": val=f"4111 1111 1111 {1000+i:04d}"
            elif kind=="dob": val="15 August 1988"
            elif kind=="ip": val=f"192.0.2.{10+i}"
            self.mapping[key]=val
        self.stats[kind]+=1
        return self.mapping[key]

    def redact_text(self,text):
        matches=[]
        for kind,pat in [("email",EMAIL_RE),("ssn",SSN_RE),("credit_card",CC_RE),("ip",IP_RE),("dob",DOB_RE)]:
            matches += [(kind,m.start(),m.end(),m.group()) for m in pat.finditer(text)]
        matches += [("phone",m.start(),m.end(),m.group()) for m in PHONE_RE.finditer(text) if valid_phone(m.group())]

        # Exact known addresses. Do not include labels such as "Registered Office:".
        for addr in sorted(KNOWN_ADDRESSES,key=len,reverse=True):
            for m in re.finditer(re.escape(addr),text,re.I): matches.append(("address",m.start(),m.end(),m.group()))
        for m in GENERIC_ADDRESS_RE.finditer(text): matches.append(("address",m.start(),m.end(),m.group()))

        for company in sorted(COMPANIES,key=len,reverse=True):
            for m in re.finditer(r"(?<![A-Za-z])"+re.escape(company)+r"(?![A-Za-z])",text,re.I): matches.append(("company",m.start(),m.end(),m.group()))
        for person in sorted(PERSONS,key=len,reverse=True):
            for m in re.finditer(r"(?<![A-Za-z])"+re.escape(person)+r"(?![A-Za-z])",text,re.I): matches.append(("name",m.start(),m.end(),m.group()))

        priority={"email":1,"ssn":1,"credit_card":1,"ip":1,"dob":1,"phone":2,"address":3,"company":4,"name":5}
        matches.sort(key=lambda x:(x[1],priority[x[0]],-(x[2]-x[1])))
        selected=[]; last_end=-1
        for item in matches:
            if item[1]>=last_end:
                selected.append(item); last_end=item[2]
        for kind,start,end,original in reversed(selected):
            text=text[:start]+self.replacement(kind,original)+text[end:]
        return text

def process_paragraph(p,r):
    if not p.text:return
    new=r.redact_text(p.text)
    if new!=p.text:
        if p.runs:
            p.runs[0].text=new
            for run in p.runs[1:]: run.text=""
        else:p.add_run(new)

def process_table(t,r,seen=None):
    seen=set() if seen is None else seen
    for row in t.rows:
        for c in row.cells:
            if c._tc in seen:continue
            seen.add(c._tc)
            for p in c.paragraphs:process_paragraph(p,r)
            for nt in c.tables:process_table(nt,r,seen)

def redact_docx(inp,out):
    d=Document(inp); r=Redactor()
    for p in d.paragraphs:process_paragraph(p,r)
    seen=set()
    for t in d.tables:process_table(t,r,seen)
    for s in d.sections:
        for p in s.header.paragraphs:process_paragraph(p,r)
        for p in s.footer.paragraphs:process_paragraph(p,r)
    d.save(out); return r

if __name__=='__main__':
    if len(sys.argv)!=3: raise SystemExit('Usage: python pii_redaction_v2.py INPUT.docx OUTPUT.docx')
    r=redact_docx(sys.argv[1],sys.argv[2])
    for k,v in r.stats.items(): print(f'{k}: {v}')
